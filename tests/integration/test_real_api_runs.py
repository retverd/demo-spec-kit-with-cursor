"""Опциональные интеграционные прогоны против реальных API (CBR/MOEX).

По умолчанию тесты пропускаются, чтобы набор был стабильным и не зависел от сети.
Для запуска установите переменную окружения RUN_REAL_API=1.
"""

from __future__ import annotations

import os
import time
from pathlib import Path

import openpyxl
import pyarrow.parquet as pq
import pytest

from src.cli.main import main


def _is_real_api_enabled() -> bool:
    value = os.getenv("RUN_REAL_API", "")
    return value.strip().lower() in {"1", "true", "yes", "y", "on"}


RUN_REAL_API = _is_real_api_enabled()
SKIP_REASON = (
    "Требуется реальная сеть: установите RUN_REAL_API=1 для запуска этих тестов."
)


@pytest.mark.skipif(not RUN_REAL_API, reason=SKIP_REASON)
@pytest.mark.parametrize("days", [1, 7, 30])
def test_real_cbr_cli_creates_parquet(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, days: int
) -> None:
    """T022/T024: реальный прогон CBR и базовая регрессия артефактов (имя/метаданные/строки)."""
    monkeypatch.chdir(tmp_path)

    start = time.monotonic()
    exit_code = main(["cbr", "--days", str(days)])
    elapsed = time.monotonic() - start

    assert exit_code == 0

    files = list(tmp_path.glob("rub_usd_*_to_*_*.parquet"))
    assert len(files) == 1

    parquet_path = files[0]
    table = pq.read_table(parquet_path)
    assert table.num_rows == days

    metadata = pq.read_metadata(parquet_path).metadata or {}
    for key in (b"report_date", b"period_start", b"period_end", b"data_source"):
        assert key in metadata
    assert metadata[b"data_source"].decode("utf-8") == "CBR"

    throughput = days / elapsed if elapsed > 0 else float("inf")
    print(
        f"CBR real API: days={days}, elapsed={elapsed:.2f}s, throughput={throughput:.2f} rows/s"
    )


@pytest.mark.skipif(not RUN_REAL_API, reason=SKIP_REASON)
@pytest.mark.parametrize("days", [1, 7, 30])
def test_real_moex_cli_creates_xlsx(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, days: int
) -> None:
    """T023/T024: реальный прогон MOEX и базовая регрессия XLSX (лист/заголовки/строки)."""
    monkeypatch.chdir(tmp_path)

    start = time.monotonic()
    exit_code = main(["moex-lqdt", "--days", str(days)])
    elapsed = time.monotonic() - start

    assert exit_code == 0

    files = list(tmp_path.glob("lqdt_tqtf_*_to_*_*.xlsx"))
    assert len(files) == 1

    xlsx_path = files[0]
    workbook = openpyxl.load_workbook(xlsx_path)
    assert "candles" in workbook.sheetnames
    sheet = workbook["candles"]

    headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert headers == ["Date", "Open", "High", "Low", "Close", "Volume"]
    assert sheet.max_row == days + 1  # +1 header row

    throughput = days / elapsed if elapsed > 0 else float("inf")
    print(
        f"MOEX real API: days={days}, elapsed={elapsed:.2f}s, throughput={throughput:.2f} rows/s"
    )


@pytest.mark.skipif(not RUN_REAL_API, reason=SKIP_REASON)
def test_real_moex_perf_smoke(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    """T037: простой перф‑смоук MOEX (по умолчанию порог 20 секунд)."""
    monkeypatch.chdir(tmp_path)

    max_seconds = float(os.getenv("MOEX_PERF_MAX_SECONDS", "20"))
    days = int(os.getenv("MOEX_PERF_DAYS", "7"))

    start = time.monotonic()
    exit_code = main(["moex-lqdt", "--days", str(days)])
    elapsed = time.monotonic() - start

    assert exit_code == 0
    assert (
        elapsed <= max_seconds
    ), f"MOEX perf failed: {elapsed:.2f}s > {max_seconds:.2f}s"
